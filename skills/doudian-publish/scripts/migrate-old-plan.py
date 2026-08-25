#!/usr/bin/env python3
"""把旧版货盘 Excel（Sheet1 价格表 + Sheet2 SKU 命名表）转成新模板。

旧格式的问题：
  - Sheet1 的价格分「上衣价格/裤子价格/套装价格」三列，Sheet2 的部件名却是
    上衣/裤子/裙子/短裙/裤裙/短裤/七分裤/背心/丝巾 等自由文本，两边对不上；
  - 库存只按尺码给一组数，不区分部件；
  - 图片列是 WPS 的 DISPIMG 嵌入公式，读不到真实文件名；
  - Sheet2 用「列=款号、行=SKU」的转置布局，且靠空行分组，机器解析脆弱。

用法:
  python3 migrate-old-plan.py <旧货盘.xlsx> [-o 输出.xlsx] [--material-root 链接图片]
"""
import os
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("缺少 openpyxl，请先 pip3 install openpyxl")
    sys.exit(1)

SIZE_ORDER = ["S", "M", "L", "XL"]

# Sheet2 部件名 -> Sheet1 价格列
TOP_WORDS = ("上衣", "衬衫", "T恤", "背心", "吊带", "外套", "马甲")
BOTTOM_WORDS = ("裤子", "裙子", "短裙", "裤裙", "短裤", "七分裤", "半身裙", "长裤")
SET_WORDS = ("套装",)


def classify(part):
    """判断部件属于 上衣价格 / 裤子价格 / 套装价格 哪一列"""
    if any(w in part for w in SET_WORDS):
        return "套装价格"
    if any(w in part for w in BOTTOM_WORDS):
        return "裤子价格"
    if any(w in part for w in TOP_WORDS):
        return "上衣价格"
    return None


def first_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"\d+(?:\.\d+)?", str(v))
    return float(m.group()) if m else None


def all_numbers(v):
    """一个单元格里可能写了多个价（如「59.9（上衣）\n39.9（吊带）」）"""
    if v is None:
        return []
    if isinstance(v, (int, float)):
        return [float(v)]
    return [float(x) for x in re.findall(r"\d+(?:\.\d+)?", str(v))]


def base_code(raw):
    """从「GZ0018（两件套/粉色套装）」提取 GZ0018"""
    s = str(raw or "").split("\n")[0]
    m = re.match(r"[A-Za-z]*\d+[A-Za-z0-9-]*", s.strip())
    return m.group() if m else s.strip()


def read_sheet1(ws):
    """返回 {款号: {价格列: 数值, 库存: {尺码: 数}}}"""
    out = {}
    # 表头在第 2 行，尺码子表头在第 3 行
    hdr = [str(c.value or "").strip() for c in ws[2]]
    col = {}
    for i, h in enumerate(hdr):
        if h in ("款号", "上衣价格", "裤子价格", "套装价格"):
            col[h] = i
    size_start = hdr.index("尺码（套装）") if "尺码（套装）" in hdr else None

    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[col.get("款号", 1)]:
            continue
        code = base_code(row[col["款号"]])
        rec = out.setdefault(code, {"价格": {}, "库存": {}, "原始款号": [], "歧义价": {}})
        rec["原始款号"].append(str(row[col["款号"]]).replace("\n", " "))
        for k in ("上衣价格", "裤子价格", "套装价格"):
            if k not in col:
                continue
            raw = row[col[k]]
            nums = all_numbers(raw)
            if len(nums) > 1:
                # 单元格里多个价，无法自动判断哪个属于哪个部件
                rec["歧义价"][k] = str(raw).replace("\n", " / ")
            if nums:
                rec["价格"].setdefault(k, nums[0])
        if size_start is not None:
            for j, sz in enumerate(SIZE_ORDER):
                v = row[size_start + j] if size_start + j < len(row) else None
                n = first_number(v)
                if n is not None:
                    rec["库存"][sz] = int(rec["库存"].get(sz, 0) + n)
    return out


def read_sheet2(ws):
    """返回 {款号: [部件名, ...]}，保持出现顺序去重"""
    out = {}
    codes = [str(c.value).strip() if c.value else None for c in ws[1]]
    for ci, code in enumerate(codes):
        if not code:
            continue
        parts = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            v = row[ci] if ci < len(row) else None
            if not v:
                continue
            s = str(v).strip()
            # 格式 款号-部件-尺码，也可能是 GT440-B-7247-上衣-S
            seg = s.split("-")
            if len(seg) < 2:
                continue
            size = seg[-1]
            part = seg[-2]
            if size not in SIZE_ORDER and not size.isdigit() and size != "F":
                part = seg[-1]
            if part and part not in parts:
                parts.append(part)
        out[base_code(code)] = parts
    return out


def guess_material_dir(root, code):
    """旧目录名可能是 0021 / GZ0021 / TM036，都试一遍"""
    if not root or not os.path.isdir(root):
        return ""
    cands = [code, re.sub(r"^[A-Za-z]+", "", code), code.upper(), code.lower()]
    for c in cands:
        p = os.path.join(root, c)
        if os.path.isdir(p):
            return os.path.relpath(p, os.path.dirname(os.path.abspath(root)))
    return ""


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    out = "抖店上架模板-已迁移.xlsx"
    mat_root = ""
    args = sys.argv[2:]
    for i, a in enumerate(args):
        if a in ("-o", "--out") and i + 1 < len(args):
            out = args[i + 1]
        if a == "--material-root" and i + 1 < len(args):
            mat_root = args[i + 1]

    wb = openpyxl.load_workbook(src, data_only=True)
    s1 = read_sheet1(wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0])
    s2 = read_sheet2(wb["Sheet2"]) if "Sheet2" in wb.sheetnames else {}

    # 生成新模板
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    tpl = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "templates", "make-template.py")
    os.system(f'python3 "{tpl}" "{out}" >/dev/null')

    nwb = openpyxl.load_workbook(out)
    ws = nwb["商品"]
    # 清掉示例行
    ws.delete_rows(3, ws.max_row)

    headers = [str(c.value or "").rstrip("*") for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(headers)}

    notes = []
    r = 3
    for code, rec in s1.items():
        parts = s2.get(code, [])
        if not parts:
            notes.append(f"{code}: Sheet2 无 SKU 命名，规格值需人工填写")
        price_pairs, unmapped = [], []
        for p in parts:
            key = classify(p)
            if key and key in rec["价格"]:
                price_pairs.append(f"{p}:{rec['价格'][key]}")
            else:
                unmapped.append(p)
        if unmapped:
            notes.append(f"{code}: 部件 {'、'.join(unmapped)} 无法自动匹配价格，需人工补")

        for k, raw in rec.get("歧义价", {}).items():
            notes.append(f"{code}: 旧表「{k}」单元格含多个价（{raw}），"
                         f"已取第一个，其余部件价格需人工改正")

        combos = [p for p in parts if ("+" in p or "＋" in p) and "套装" not in p]
        if combos:
            notes.append(f"{code}: 「{'、'.join(combos)}」是组合件，"
                         f"自动归类的价格不可靠，请人工确认")

        stock = "|".join(f"{s}:{int(rec['库存'].get(s, 0))}" for s in SIZE_ORDER
                         if s in rec["库存"])

        vals = {
            "款号": code,
            "商品标题": "",
            "类目": "服装>女装>时尚套装",
            "面料材质": "聚酯纤维",
            "材质含量": "",
            "素材目录": guess_material_dir(mat_root, code) or f"链接图片/{code}",
            "规格组名": "颜色分类",
            "规格值": "|".join(parts),
            "尺码": "|".join(s for s in SIZE_ORDER if s in rec["库存"]) or "S|M|L|XL",
            "价格": "|".join(price_pairs),
            "库存": stock,
            "库存明细": "",
            "商家编码规则": "",
            "商品状态": "下架",
            "尺码表模板": "身高体重",
            "运费模板": "偏远区域限购",
            "备注": " / ".join(rec["原始款号"]),
        }
        for h, c in idx.items():
            ws.cell(row=r, column=c, value=vals.get(h, ""))
        r += 1

    # 迁移说明 Sheet
    ws4 = nwb.create_sheet("迁移说明")
    ws4.column_dimensions["A"].width = 100
    lines = [
        "本文件由旧版货盘自动迁移生成，以下内容必须人工复核后才能用于上架：",
        "",
        "1. 类目：统一填了「服装>女装>时尚套装」，非套装款需改成实际类目",
        "   （如纯上衣款 → 服装>女装>T恤/Polo衫）。",
        "2. 面料材质：统一填了「聚酯纤维」，需按实际吊牌修正。",
        "3. 素材目录：按款号猜的路径，需核对是否存在；且目录内文件要按",
        "   「素材命名规范」Sheet 重命名（主图01.jpg / 详情图01.jpg / SKU_上衣.jpg）。",
        "4. 库存：旧表只按尺码给一组数，未区分部件，迁移后所有部件共用同一库存。",
        "   若各部件库存不同，请填「库存明细」列。",
        "5. 一款多色（如 GZ0018 粉色/黑色）在旧表里是多行，迁移后合并成一行，",
        "   规格值里含所有颜色部件，库存已按行累加，需核对。",
        "6. 旧表的图片是 WPS DISPIMG 嵌入公式，无法读出文件名，未迁移。",
        "",
        "自动检测到的问题：",
    ]
    lines.extend(f"  - {n}" for n in notes) if notes else lines.append("  （无）")
    for i, t in enumerate(lines, start=1):
        c = ws4.cell(row=i, column=1, value=t)
        if i == 1 or t.endswith("："):
            c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    nwb.save(out)
    print(f"已生成 {out}（{r - 3} 款）")
    if notes:
        print("需人工处理：")
        for n in notes:
            print(f"  - {n}")


if __name__ == "__main__":
    main()
