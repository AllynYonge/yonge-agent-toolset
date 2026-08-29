#!/usr/bin/env python3
"""生成抖店自动上架用的 Excel 模板（含填写说明与示例行）。

用法:
  python3 make-template.py [输出路径]
默认输出 ./抖店上架模板.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "抖店上架模板.xlsx"

# (列名, 宽度, 是否必填, 说明)
COLS = [
    ("款号",        14, True,  "货品编号，同时写入抖店「货号」字段，也是素材目录名。例：GZ0021"),
    ("参考图",       30, True,  "直接把图片插入本列单元格（可多张）。必须用「浮动图片」插入，不能用 Excel 365 的「单元格内图片」；也可改填图片路径或目录，多张用 | 分隔"),
    ("商品标题",     46, False, "留空则由 AI 依据款号/类目/主图生成 15-60 字符（8-30 汉字）标题"),
    ("类目",        22, True,  "抖店三级类目全路径，用 > 分隔。例：服装>女装>时尚套装"),
    ("面料材质",     16, True,  "单一材质名，脚本自动补 100%。例：聚酯纤维"),
    ("材质含量",      9, False, "留空默认 100，多材质时用 分号 分隔：聚酯纤维:90;氨纶:10"),
    ("素材目录",     26, True,  "该款素材所在目录（相对 --root 或绝对路径）。例：链接图片/0021"),
    ("规格组名",     12, True,  "第一组规格的名称，服饰通常为 颜色分类"),
    ("规格值",       30, True,  "第一组规格值，用 | 分隔，顺序即前台展示顺序。例：上衣|裙子|套装"),
    ("尺码",        22, True,  "第二组规格值，用 | 分隔。例：S|M|L|XL"),
    ("价格",        34, True,  "按「规格值:价格」写，用 | 分隔，顺序须与规格值一致。例：上衣:49.9|裙子:59.9|套装:99.9"),
    ("库存",        34, True,  "按「尺码:库存」写，用 | 分隔，全规格共用。例：S:360|M:350|L:330|XL:110"),
    ("库存明细",     40, False, "需要按规格分别设库存时填，优先级高于「库存」。格式：上衣-S:100|上衣-M:80|套装-S:60"),
    ("商家编码规则",  26, False, "留空默认 {款号}-{规格值}-{尺码}，如 GZ0021-上衣-S"),
    ("商品状态",     10, False, "上架 / 下架，留空取全局默认"),
    ("尺码表模板",   16, False, "抖店已存尺码表模板名，如 身高体重。留空则不套用模板"),
    ("运费模板",     16, False, "抖店运费模板名，如 偏远区域限购 / 包邮。留空取全局默认"),
    ("备注",        30, False, "不写入抖店，仅供人工记录"),
]

EXAMPLES = [
    ["GZ0021", "", "", "服装>女装>时尚套装", "聚酯纤维", "", "链接图片/0021",
     "颜色分类", "上衣|裙子|套装", "S|M|L|XL",
     "上衣:49.9|裙子:59.9|套装:99.9", "S:360|M:350|L:330|XL:110", "",
     "", "下架", "身高体重", "偏远区域限购", "示例行，正式使用时请删除"],
    ["GZ0018", "", "", "服装>女装>时尚套装", "聚酯纤维", "", "链接图片/0018",
     "颜色分类", "粉色上衣|黑色上衣|裙子|粉色套装|黑色套装", "S|M|L|XL",
     "粉色上衣:49.9|黑色上衣:49.9|裙子:59.9|粉色套装:99.9|黑色套装:99.9",
     "S:90|M:14|L:60|XL:50", "",
     "", "下架", "身高体重", "偏远区域限购", "示例行（一款多色，颜色分类里区分色系），正式使用时请删除"],
]

wb = Workbook()

# ---------- Sheet 商品 ----------
ws = wb.active
ws.title = "商品"

hdr_req = PatternFill("solid", fgColor="FFD966")   # 必填 橙
hdr_opt = PatternFill("solid", fgColor="D9E1F2")   # 选填 蓝
bold = Font(bold=True, size=11)

for c, (name, width, required, note) in enumerate(COLS, start=1):
    cell = ws.cell(row=1, column=c, value=name + ("*" if required else ""))
    cell.font = bold
    cell.fill = hdr_req if required else hdr_opt
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.comment = None
    ws.column_dimensions[get_column_letter(c)].width = width
    # 第 2 行写说明
    n = ws.cell(row=2, column=c, value=note)
    n.font = Font(size=9, color="808080")
    n.alignment = Alignment(wrap_text=True, vertical="top")

ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 58

for r, row in enumerate(EXAMPLES, start=3):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)

# 数据行拉高，便于在「参考图」列里插图后还能看清其他列
for r in range(3, 3 + max(len(EXAMPLES), 20)):
    ws.row_dimensions[r].height = 96

ws.freeze_panes = "A3"

# 「商品状态」列位置随列定义变化，动态取列字母，避免加列后校验落到别的列上
status_col = get_column_letter(next(i for i, c in enumerate(COLS, start=1) if c[0] == "商品状态"))
dv = DataValidation(type="list", formula1='"上架,下架"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"{status_col}3:{status_col}500")

# ---------- Sheet 素材命名规范 ----------
ws2 = wb.create_sheet("素材命名规范")
rules = [
    ["文件名规范", "用途", "抖店字段", "数量要求"],
    ["主图01.jpg … 主图05.jpg", "商品主图，01 为首图", "主图（1:1）", "3-5 张，必填"],
    ["详情图01.jpg … 详情图NN.jpg", "商品详情长图", "商品详情", "≥1 张，建议 5-10 张"],
    ["SKU_<规格值>.jpg", "第一组规格的规格图", "规格图", "每个规格值 1 张，必填"],
    ["尺码表.jpg", "尺码表图（放在详情图末尾）", "商品详情", "选填但强烈建议"],
    [],
    ["说明", "", "", ""],
    ["1. 文件名严格按上表写，脚本按前缀识别用途，不做模糊推断。", "", "", ""],
    ["2. SKU_ 后面必须与「规格值」列中的值完全一致，如 SKU_上衣.jpg / SKU_粉色套装.jpg。", "", "", ""],
    ["3. 扩展名支持 jpg / jpeg / png。jfif 需先转换（脚本可自动转）。", "", "", ""],
    ["4. 主图要求 ≥600x600px，详情图宽度 620-1290px，单张 <5MB。", "", "", ""],
    ["5. 目录名建议与款号一致，便于核对。", "", "", ""],
    ["6. 「参考图」不走文件名规范：图片直接插进「商品」Sheet 的「参考图」单元格，"
     "脚本会导出到 <素材目录>/参考图/ 供人工查看，不上传抖店。", "", "", ""],
]
for r, row in enumerate(rules, start=1):
    for c, v in enumerate(row, start=1):
        cell = ws2.cell(row=r, column=c, value=v)
        if r == 1 or v == "说明":
            cell.font = bold
            if r == 1:
                cell.fill = hdr_opt
for c, w in enumerate([32, 26, 16, 22], start=1):
    ws2.column_dimensions[get_column_letter(c)].width = w

# ---------- Sheet 全局默认 ----------
ws3 = wb.create_sheet("全局默认")
defaults = [
    ["配置项", "值", "说明"],
    ["品牌", "对门的猫", "抖店「品牌」必填项，无品牌时填 无品牌"],
    ["发货模式", "现货", "现货 / 现货预售混合。目前取页面默认值（实测默认就是现货）"],
    ["发货时效", "48小时", "当日发 / 次日发 / 48小时，写入页面「现货发货时间」"],
    ["运费模板", "偏远区域限购", "抖店后台已存在的运费模板名"],
    ["售后政策", "7天无理由退货", "发布页无独立控件，取页面默认值"],
    ["订单库存计数", "付款减库存", "下单减库存 / 付款减库存"],
    ["商品状态", "下架", "批量上架建议先设 下架，人工复核后再统一上架"],
    ["尺码表模板", "身高体重", "抖店后台已存的尺码表模板名"],
    ["商家编码规则", "{款号}-{规格值}-{尺码}", "支持占位符 {款号} {规格值} {尺码}"],
]
for r, row in enumerate(defaults, start=1):
    for c, v in enumerate(row, start=1):
        cell = ws3.cell(row=r, column=c, value=v)
        if r == 1:
            cell.font = bold
            cell.fill = hdr_opt
for c, w in enumerate([20, 30, 48], start=1):
    ws3.column_dimensions[get_column_letter(c)].width = w

wb.save(OUT)
print(f"已生成 {OUT}")
